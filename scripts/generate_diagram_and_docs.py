"""
Generate Architecture Block Diagram, PDF Document, and Word (.docx) Document
=============================================================================
Creates:
  1. backend/static/downloads/architecture_block_diagram.png (High-res 300 DPI vector-styled diagram)
  2. backend/static/downloads/Bone_Fracture_System_Architecture.pdf (Formal ReportLab PDF)
  3. backend/static/downloads/Bone_Fracture_System_Architecture.docx (Professional Word doc with diagram & tables)
  4. Copies to workspace root for direct access
"""

import os
import sys
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as patches

# ---------------------------------------------------------------------------
# 1. GENERATE HIGH-RESOLUTION SYSTEM ARCHITECTURE BLOCK DIAGRAM
# ---------------------------------------------------------------------------
def generate_block_diagram_image(output_path: str):
    fig = plt.figure(figsize=(16, 11), dpi=300)
    ax = fig.add_subplot(111)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 11)
    ax.axis('off')

    # Color palette
    c_bg = "#f8fafc"
    c_card_bg = "#ffffff"
    c_primary = "#1e3a8a"     # Navy
    c_blue = "#3b82f6"        # Blue
    c_teal = "#0d9488"        # Teal
    c_purple = "#7c3aed"      # Purple
    c_red = "#dc2626"         # Red
    c_green = "#16a34a"       # Green
    c_gray = "#64748b"        # Gray text
    c_border = "#cbd5e1"

    fig.patch.set_facecolor(c_bg)
    ax.set_facecolor(c_bg)

    # Title Banner
    title_box = patches.FancyBboxPatch(
        (0.8, 9.8), 14.4, 0.95,
        boxstyle="round,pad=0.1,rounding_size=0.15",
        facecolor=c_primary, edgecolor="none", zorder=2
    )
    ax.add_patch(title_box)
    ax.text(8.0, 10.4, "END-TO-END DEEP LEARNING BONE FRACTURE RADIOGRAPH ANALYSIS SYSTEM",
            color="white", fontsize=15, fontweight="bold", ha="center", va="center", zorder=3)
    ax.text(8.0, 10.0, "System Architecture: Multi-Task Classification, Spatial Localization, Disciplined Captioning & Web Dashboard",
            color="#93c5fd", fontsize=10, ha="center", va="center", zorder=3)

    # Helper function for drawing component boxes
    def draw_box(x, y, w, h, title, subtitle, header_color, details=[], icon=""):
        # Shadow
        shadow = patches.FancyBboxPatch(
            (x + 0.05, y - 0.05), w, h,
            boxstyle="round,pad=0.08,rounding_size=0.12",
            facecolor="#e2e8f0", edgecolor="none", zorder=1
        )
        ax.add_patch(shadow)

        # Main Card
        card = patches.FancyBboxPatch(
            (x, y), w, h,
            boxstyle="round,pad=0.08,rounding_size=0.12",
            facecolor="white", edgecolor=c_border, linewidth=1.2, zorder=2
        )
        ax.add_patch(card)

        # Header Bar
        header = patches.FancyBboxPatch(
            (x, y + h - 0.55), w, 0.55,
            boxstyle="round,pad=0.08,rounding_size=0.1",
            facecolor=header_color, edgecolor="none", zorder=3
        )
        ax.add_patch(header)
        ax.text(x + w/2, y + h - 0.28, title,
                color="white", fontsize=10, fontweight="bold", ha="center", va="center", zorder=4)

        if subtitle:
            ax.text(x + w/2, y + h - 0.75, subtitle,
                    color=c_gray, fontsize=8, fontweight="bold", ha="center", va="center", zorder=4)

        # Bullet details
        y_text = y + h - 1.05
        for line in details:
            ax.text(x + 0.2, y_text, f"- {line}",
                    color="#334155", fontsize=7.8, ha="left", va="center", zorder=4)
            y_text -= 0.32

    # Helper function for drawing arrows
    def draw_arrow(x1, y1, x2, y2, label="", color="#475569"):
        ax.annotate(
            "", xy=(x2, y2), xytext=(x1, y1),
            arrowprops=dict(
                arrowstyle="-|>", color=color, lw=1.8,
                mutation_scale=14, shrinkA=3, shrinkB=3
            ), zorder=5
        )
        if label:
            mx, my = (x1 + x2) / 2, (y1 + y2) / 2
            ax.text(mx, my + 0.12, label, color=color, fontsize=7.5,
                    fontweight="bold", ha="center", va="center",
                    bbox=dict(boxstyle="round,pad=0.2", facecolor="white", edgecolor=color, lw=0.8),
                    zorder=6)

    # ------------------ STAGE 1: INPUT & PREPROCESSING ------------------
    draw_box(0.8, 6.7, 3.2, 2.7, "INPUT & PREPROCESSING", "Data Preparation Module", c_teal, [
        "Plain radiograph (PNG/JPG/BMP)",
        "In-memory stream validation",
        "Grayscale conversion ('L')",
        "Aspect-ratio padding -> 224x224",
        "Standardization & Tensor Scale",
        "Mean=[0.5]*3, Std=[0.5]*3"
    ])

    # ------------------ STAGE 2: TIER 1 SHARED BACKBONE ------------------
    draw_box(4.8, 6.7, 3.4, 2.7, "TIER 1: RESNET-50 BACKBONE", "Multi-Task Feature Extractor", c_blue, [
        "Pretrained CNN Feature Extractor",
        "conv1 -> bn1 -> relu -> maxpool",
        "layer1 (3 blocks) -> layer2 (4 blk)",
        "layer3 (6 blocks) -> layer4 (3 blk)",
        "Adaptive Average Pooling",
        "2,048-dim Latent Feature Vector"
    ])

    # ------------------ STAGE 3: HEADS ------------------
    draw_box(9.0, 7.6, 3.1, 1.8, "ANATOMICAL HEAD", "7-Class Softmax Classifier", "#2563eb", [
        "Linear(2048, 512) -> ReLU -> Drop(0.3)",
        "Linear(512, 7) -> CrossEntropy",
        "Arm, Foot, Hand, Lower leg,",
        "Thigh, pelvis, wrist"
    ])

    draw_box(9.0, 5.5, 3.1, 1.8, "FRACTURE HEAD", "Binary BCE Classifier", c_red, [
        "Linear(2048, 256) -> ReLU -> Drop(0.3)",
        "Linear(256, 1) -> Sigmoid",
        "Weighted pos_weight = 2.438",
        "Decision Threshold >= 0.50"
    ])

    # ------------------ STAGE 4: TIER 2 LOCALIZATION ------------------
    draw_box(12.7, 5.5, 2.7, 3.9, "TIER 2: LOCALIZATION", "Faster R-CNN Detector", "#b91c1c", [
        "Target Size: 384x384 (RGB)",
        "Backbone: MobileNetV3-Large",
        "Feature Pyramid Network (FPN)",
        "Region Proposal Network (RPN)",
        "RoI Align Box Predictor",
        "Calibrated Thresh >= 0.10",
        "Box scaling to native image",
        "Outputs [x1, y1, x2, y2]"
    ])

    # ------------------ STAGE 5: FACTUAL CAPTIONING ------------------
    draw_box(9.0, 3.2, 6.4, 1.9, "TIER 3: FACTUAL CAPTION GENERATOR", "Zero-Hallucination Medical Synthesizer", c_purple, [
        "Rule-based deterministic synthesizer mapping validated CV outputs",
        "No LLM hallucination: never invents age, sex, trauma mechanism, severity, or treatment",
        "Template: 'X-ray of the {region} showing a fracture.' / '...with no fracture detected.'",
        "Appends localization details when spatial lesions are confirmed"
    ])

    # ------------------ STAGE 6: REST API BACKEND ------------------
    draw_box(0.8, 0.7, 6.8, 2.0, "FASTAPI REST BACKEND (Port 8000)", "High-Throughput Diagnostic Microservice", "#0f766e", [
        "Singleton InferenceService: loads PyTorch models once into memory",
        "Endpoints: POST /predict (multipart/form-data), GET /health, GET /visualizations/{id}",
        "Safe in-memory image buffer (Zero temporary files on disk)",
        "Strict Pydantic response schema + Path-traversal attack protection"
    ])

    # ------------------ STAGE 7: STREAMLIT FRONTEND ------------------
    draw_box(8.5, 0.7, 6.9, 2.0, "STREAMLIT DASHBOARD UI (Port 8501)", "Interactive Diagnostic Web Interface", "#c2410c", [
        "Drag-and-Drop Radiograph Upload (PNG, JPG, BMP, WebP)",
        "Three-Card Diagnostic Row: Anatomical Region, Fracture Status, Lesion Count",
        "Visual Diagnostic Overlay: Original X-ray with bold red bounding box callouts",
        "Interactive Lesion Coordinate Table + Raw JSON response inspector"
    ])

    # Connectors
    draw_arrow(4.0, 8.0, 4.8, 8.0, "224x224x3")
    draw_arrow(8.2, 8.3, 9.0, 8.3, "Features")
    draw_arrow(8.2, 7.2, 9.0, 6.6, "Features")
    draw_arrow(12.1, 6.4, 12.7, 6.4, "If Frac=TRUE")
    draw_arrow(10.5, 5.5, 10.5, 5.1, "")
    draw_arrow(14.0, 5.5, 14.0, 5.1, "")
    draw_arrow(12.2, 5.1, 12.2, 4.2, "Structured Predictions")
    draw_arrow(2.4, 6.7, 2.4, 2.7, "API Client")
    draw_arrow(7.6, 1.7, 8.5, 1.7, "JSON Payload")

    # Save
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    plt.savefig(output_path, dpi=300, bbox_inches='tight', facecolor=fig.get_facecolor())
    plt.close()
    print(f"[OK] Saved block diagram image to: {output_path}")

# ---------------------------------------------------------------------------
# 2. GENERATE FORMAL PDF DOCUMENT (Using ReportLab)
# ---------------------------------------------------------------------------
def generate_pdf_document(pdf_path: str, diagram_img_path: str):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, HRFlowable
    )

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e3a8a'),
        fontName='Helvetica-Bold',
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#475569'),
        spaceAfter=15
    )

    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1e3a8a'),
        fontName='Helvetica-Bold',
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1e293b')
    )

    story = []

    # Title & Metadata
    story.append(Paragraph("Bone Fracture Detection & Captioning System", title_style))
    story.append(Paragraph("System Architecture Specification & Technical Block Diagram Document | Version 1.0", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=12))

    # Section 1: Overview
    story.append(Paragraph("1. System Architectural Overview", h2_style))
    overview_text = (
        "The Bone Fracture Image Captioning and Detection platform provides automated, multi-task radiograph analysis "
        "consisting of three primary execution tiers: (1) Multi-Task Shared Convolutional Backbone for 7-class anatomical "
        "site classification and binary fracture detection, (2) Faster R-CNN with Feature Pyramid Network (FPN) for spatial "
        "lesion localization, and (3) Disciplined Factual Caption Synthesizer generating non-hallucinatory medical text."
    )
    story.append(Paragraph(overview_text, body_style))
    story.append(Spacer(1, 10))

    # Embed High-Resolution Block Diagram
    if os.path.exists(diagram_img_path):
        story.append(Paragraph("2. System Architecture Block Diagram", h2_style))
        # Letter page usable width = 612 - 80 = 532 pt
        story.append(RLImage(diagram_img_path, width=532, height=365))
        story.append(Spacer(1, 10))

    # Section 3: Component Specification Table
    story.append(Paragraph("3. Deep Learning Component Specifications", h2_style))
    table_data = [
        ["Component", "Architecture", "Loss / Objective", "Output Specification"],
        ["Classifier Backbone", "ResNet-50 (Shared)", "Two-stage Transfer Learning", "2,048-dim feature representation"],
        ["Anatomical Head", "2-Layer MLP + Drop(0.3)", "Cross-Entropy Loss", "7 classes: Arm, Foot, Hand, Lower leg, Thigh, pelvis, wrist"],
        ["Fracture Head", "2-Layer MLP + Drop(0.3)", "BCEWithLogits (pos_weight=2.438)", "Binary fracture probability (Threshold >= 0.50)"],
        ["Spatial Detector", "Faster R-CNN (MobileNetV3 FPN)", "RPN + RoI Box Regression", "Bounding boxes [x1, y1, x2, y2] (Thresh >= 0.10)"],
        ["Caption Generator", "Template Synthesizer", "Deterministic rule engine", "Factual natural language description"]
    ]
    t = Table(table_data, colWidths=[110, 140, 132, 150])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    # Section 4: Workflow
    story.append(Paragraph("4. End-to-End Dashboard Interaction Workflow", h2_style))
    workflow_text = (
        "<b>Step 1: Upload</b> — Doctor or clinical user uploads a plain radiograph into the Streamlit dashboard.<br/>"
        "<b>Step 2: Analysis Trigger</b> — The 'Analyze X-ray' button dispatches an asynchronous multipart request to FastAPI <code>POST /predict</code>.<br/>"
        "<b>Step 3: Multi-Task Classification</b> — ResNet-50 predicts the anatomical site and fracture probability.<br/>"
        "<b>Step 4: Spatial Localization</b> — If fracture is positive, Faster R-CNN detects lesion bounds and generates a visual diagnostic overlay.<br/>"
        "<b>Step 5: Visual Presentation</b> — Dashboard renders three KPI cards, factual caption, bounding box overlay, and coordinate tables."
    )
    story.append(Paragraph(workflow_text, body_style))

    doc.build(story)
    print(f"[OK] Saved PDF document to: {pdf_path}")

# ---------------------------------------------------------------------------
# 3. GENERATE FORMAL WORD (.DOCX) DOCUMENT (Using python-docx)
# ---------------------------------------------------------------------------
def generate_docx_document(docx_path: str, diagram_img_path: str):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = docx.Document()

    # Set standard margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Helper for cell shading
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        tcPr.append(shd)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Bone Fracture Detection & Captioning System")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("System Architecture Specification & Technical Block Diagram Document")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    # Section 1
    h1 = doc.add_heading("1. Executive Summary & Architectural Overview", level=1)
    h1.paragraph_format.space_before = Pt(10)
    p_desc = doc.add_paragraph(
        "This engineering specification details the deep-learning bone fracture analysis platform. "
        "The solution combines multi-task deep convolutional representation learning (ResNet-50), "
        "dedicated spatial lesion localization (Faster R-CNN with MobileNetV3 Feature Pyramid Network), "
        "a zero-hallucination factual clinical captioning synthesizer, a high-throughput FastAPI REST microservice, "
        "and an interactive Streamlit diagnostic web dashboard."
    )
    p_desc.paragraph_format.line_spacing = 1.15

    # Section 2: Embed Block Diagram Image
    doc.add_heading("2. System Architecture Block Diagram", level=1)
    if os.path.exists(diagram_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(10)
        doc.add_picture(diagram_img_path, width=Inches(6.8))

    # Section 3: Architecture Specification Table
    doc.add_heading("3. Component Specifications & Hyperparameters", level=1)
    table = doc.add_table(rows=6, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Component", "Base Architecture", "Loss / Objective Function", "Output Specification"]
    row_data = [
        ["Classifier Backbone", "ResNet-50 (Stage 1 frozen, Stage 2 fine-tuned)", "Transfer learning feature extractor", "2,048-dim feature representation"],
        ["Anatomical Head", "2-Layer MLP (2048->512->7) + Dropout(0.3)", "Multi-class Cross-Entropy Loss", "7 classes (Arm, Foot, Hand, Lower leg, Thigh, pelvis, wrist)"],
        ["Fracture Head", "2-Layer MLP (2048->256->1) + Dropout(0.3)", "BCEWithLogitsLoss (pos_weight = 2.438)", "Fracture probability (Decision Threshold >= 0.50)"],
        ["Spatial Detector", "Faster R-CNN (MobileNetV3-Large FPN)", "RPN Loss + Fast R-CNN Box Regression", "Bounding boxes [x1, y1, x2, y2] (Threshold >= 0.10)"],
        ["Caption Synthesizer", "Deterministic rule engine template mapper", "Structured output binding", "Disciplined natural-language clinical description"]
    ]

    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E3A8A")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)

    # Format Data Rows
    for r_idx, row in enumerate(row_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

    # Section 4: Dashboard Operation Flow
    doc.add_heading("4. How the Dashboard Works", level=1)
    flow_points = [
        ("1. Image Upload", "Doctor or radiologist uploads plain bone radiograph (PNG, JPG, BMP) into the drag-and-drop file uploader."),
        ("2. Real-Time In-Memory Dispatch", "Upon clicking 'Analyze X-ray', Streamlit sends raw bytes to FastAPI POST /predict without temporary file disk writes."),
        ("3. Multi-Task Deep Inference", "ResNet-50 simultaneously computes 7-class anatomical probability and weighted binary fracture probability."),
        ("4. Spatial Localization & Overlay", "When fracture is confirmed, Faster R-CNN identifies cortical fracture borders and generates an annotated diagnostic visual overlay."),
        ("5. Factual Caption Synthesis", "The disciplined captioning engine synthesizes an explicit, non-hallucinatory description."),
        ("6. KPI Display & Coordinate Inspection", "Dashboard presents three color-coded metric cards, natural language text, visual image overlay, and exact bounding box coordinate table.")
    ]

    for title, desc in flow_points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.space_before = Pt(3)
        p_pt.paragraph_format.space_after = Pt(3)
        r_b = p_pt.add_run(f"• {title}: ")
        r_b.bold = True
        r_b.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
        r_d = p_pt.add_run(desc)
        r_d.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.save(docx_path)
    print(f"[OK] Saved Word (.docx) document to: {docx_path}")

# ---------------------------------------------------------------------------
# MAIN EXECUTION
# ---------------------------------------------------------------------------
if __name__ == '__main__':
    downloads_dir = "backend/static/downloads"
    os.makedirs(downloads_dir, exist_ok=True)
    os.makedirs("reports", exist_ok=True)

    diagram_png = os.path.join(downloads_dir, "architecture_block_diagram.png")
    pdf_out = os.path.join(downloads_dir, "Bone_Fracture_System_Architecture.pdf")
    docx_out = os.path.join(downloads_dir, "Bone_Fracture_System_Architecture.docx")

    # 1. Generate Diagram Image
    generate_block_diagram_image(diagram_png)
    # Also save a copy in reports/
    generate_block_diagram_image("reports/architecture_block_diagram.png")

    # 2. Generate PDF Document
    generate_pdf_document(pdf_out, diagram_png)
    # Copy to workspace root
    generate_pdf_document("Bone_Fracture_System_Architecture.pdf", diagram_png)

    # 3. Generate Word Document
    generate_docx_document(docx_out, diagram_png)
    # Copy to workspace root
    generate_docx_document("Bone_Fracture_System_Architecture.docx", diagram_png)

    print("\nALL DOCUMENTS GENERATED SUCCESSFULLY!")
